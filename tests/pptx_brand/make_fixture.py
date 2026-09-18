"""Build a small branded template (.potx), an icon folder and a vector icon deck for the pptx-brand tests.

Everything is generated, so no binary fixtures are committed.
"""

import io
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

THEME_COLORS = {"1F497D": "0B1F3A", "EEECE1": "E8EEF4", "4F81BD": "0067B8", "C0504D": "00A3A1", "9BBB59": "F2A900",
                "8064A2": "D9480F", "4BACC6": "5C2D91", "F79646": "6B7785"}
DECK_CT = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
TEMPLATE_CT = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"


def png(draw_fn, size=256, color=(11, 31, 58, 255)):
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw_fn(ImageDraw.Draw(img), color)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def make_template(path):
    prs = Presentation()
    master = prs.slide_masters[0]
    theme = master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode()
    for old, new in THEME_COLORS.items():
        xml = xml.replace(f'val="{old}"', f'val="{new}"')
    xml = xml.replace('<a:latin typeface="Calibri"/>', '<a:latin typeface="Georgia"/>', 1)  # major font only
    theme._blob = xml.encode()

    # Logo on the master, bottom-left: decoration the profiler must report as a reserved region.
    logo = png(lambda d, c: (d.rectangle([0, 80, 255, 175], fill=(0, 103, 184, 255)),
                             d.ellipse([20, 96, 84, 160], fill=(255, 255, 255, 255))))
    scratch = prs.slides.add_slide(prs.slide_layouts[6])
    pic = scratch.shapes.add_picture(io.BytesIO(logo), Inches(0.4), Inches(6.8), Inches(1.4), Inches(0.5))
    pic.name = "acme-logo"
    _, rid = master.part.get_or_add_image_part(io.BytesIO(logo))
    pic._element.xpath(".//a:blip")[0].set(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    master.shapes._spTree.append(pic._element)
    sld_id = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)

    buf = io.BytesIO()
    prs.save(buf)
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(DECK_CT, TEMPLATE_CT)
            zout.writestr(item, data)


def make_icons(directory):
    directory.mkdir(parents=True, exist_ok=True)
    shapes = {
        "shield-check": lambda d, c: d.polygon([(128, 16), (232, 56), (208, 200), (128, 244), (48, 200), (24, 56)], fill=c),
        "growth-chart": lambda d, c: [d.rectangle([40 + i * 64, 200 - i * 50, 84 + i * 64, 240], fill=c) for i in range(3)],
        "cost-coin": lambda d, c: d.ellipse([28, 28, 228, 228], fill=c),
    }
    for name, fn in shapes.items():
        (directory / f"{name}.png").write_bytes(png(fn))


def make_icon_deck(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (kind, name) in enumerate([(MSO_SHAPE.GEAR_6, "gear settings"), (MSO_SHAPE.CLOUD, "cloud hosting")]):
        shape = slide.shapes.add_shape(kind, Inches(1 + i * 3), Inches(1), Inches(1.5), Inches(1.5))
        shape.name = name
    prs.save(path)


def build(out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    make_template(out_dir / "acme-corporate.potx")
    make_icons(out_dir / "icons")
    make_icon_deck(out_dir / "icon-library.pptx")
    # Not part of any catalog: the QA test embeds it to prove stray images are flagged.
    (out_dir / "stray.png").write_bytes(png(lambda d, c: d.rectangle([10, 10, 200, 200], fill=(255, 0, 255, 255))))
    return out_dir


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1]))
