#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5", "defusedxml"]
# ///
"""End-to-end test of the pptx-brand skill: onboard -> build -> QA -> QA catches injected defects.

    uv run tests/pptx_brand/run_e2e.py [--render]

Without --render nothing needs LibreOffice, so this runs in CI.
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parents[1] / "plugins/doc-template-skills/skills/pptx-brand/scripts"
sys.path.insert(0, str(HERE))

import make_fixture  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches  # noqa: E402


def run(script, *args, expect=0):
    proc = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, text=True)
    assert proc.returncode == expect, f"{script} exited {proc.returncode}, expected {expect}:\n{proc.stdout}\n{proc.stderr}"
    return proc.stdout


def main():
    render = "--render" in sys.argv
    work = Path(tempfile.mkdtemp(prefix="pptx-brand-e2e-"))
    os.environ["PPTX_BRAND_HOME"] = str(work / "home")
    fx = make_fixture.build(work / "fixture")
    no_render = [] if render else ["--no-render"]

    # Onboard a .potx with an icon folder and a vector icon deck.
    run("onboard_template.py", fx / "acme-corporate.potx", "--icons", fx / "icons", "--icon-deck", fx / "icon-library.pptx",
        "--default", *no_render)
    tdir = work / "home/templates/acme-corporate"
    brand = json.loads((tdir / "brand.json").read_text())
    assert brand["colors"]["accent1"] == "0067B8" and brand["colors"]["dk2"] == "0B1F3A", brand["colors"]
    assert len(brand["colors"]) == 12 and all(len(v) == 6 for v in brand["colors"].values())
    assert brand["fonts"] == {"major_latin": "Georgia", "minor_latin": "Calibri"}, brand["fonts"]
    assert brand["font_safety"] == {"major_qa_reliable": False, "minor_qa_reliable": True, "fallback_body": "Calibri"}

    layouts = json.loads((tdir / "layouts.json").read_text())
    classes = {l["class"] for l in layouts["layouts"]}
    assert {"title", "section", "one_column", "two_column", "comparison", "title_only", "blank"} <= classes, classes
    assert "three_column" in layouts["coverage"]["missing_classes"]
    two_col = next(l for l in layouts["layouts"] if l["name"] == "Two Content")
    assert all("box_in" in p and "capacity_lines" in p for p in two_col["placeholders"] if p["type"] == "obj")
    assert two_col["reserved_regions_in"], "logo on the master should be a reserved region"

    assets = json.loads((tdir / "assets.json").read_text())["assets"]
    labels = {a["label"] for a in assets}
    assert {"shield check", "growth chart", "cost coin"} <= labels, labels
    assert any(a["source"] == "template" and a["used_by"] == ["slideMaster1.xml"] for a in assets), "logo not harvested"
    assert next(a for a in assets if a["label"] == "shield check")["recolorable"]
    if render:
        assert any(a["format"] == "shape" for a in assets), "vector icons from the icon deck not harvested"
        assert list((tdir / "thumbs").glob("layouts-*.jpg")) and list((tdir / "thumbs").glob("icons-*.jpg"))
    assert len((tdir / "catalog.md").read_text()) // 4 < 6000

    # Labels written by the onboarder survive a re-run (hash-keyed cache).
    for a in assets:
        if a["label"] == "cost coin":
            a.update(label="coin", tags=["cost", "budget", "coin"], labeled=True)
    (tdir / "assets.json").write_text(json.dumps({"assets": assets}))
    run("onboard_template.py", fx / "acme-corporate.potx", "--default", *no_render)
    assets = json.loads((tdir / "assets.json").read_text())["assets"]
    assert any(a["label"] == "coin" and a["labeled"] for a in assets), "label cache lost"
    assert any(a["label"] == "shield check" for a in assets), "folder assets lost on re-run without --icons"

    # A second template; the default is chosen when none is named.
    plain = work / "plain.pptx"
    Presentation().save(plain)
    run("onboard_template.py", plain, "--no-render")

    deck = work / "deck.pptx"
    out = run("fill_layout.py", HERE / "outline.md", "-o", deck)
    build = json.loads(Path(str(deck) + ".build.json").read_text())
    assert build["template"] == "acme-corporate", build["template"]
    titles = [s["title"] for s in build["slides"]]
    assert "Priorities for Q4 (cont.)" in titles, f"long list was not split:\n{out}"
    assert not any(s["unresolved_icons"] for s in build["slides"]), build["slides"]
    highlights = next(s for s in build["slides"] if s["title"] == "Highlights")
    assert len(highlights["assets"]) == 2, highlights

    prs = Presentation(str(deck))
    assert "Brand catalog assets" in prs.slides[2].notes_slide.notes_text_frame.text
    assert sum(1 for sh in prs.slides[4].shapes if sh.has_chart) == 1
    print(run("brand_qa.py", deck))

    if render:  # vector icons harvested from the icon deck are cloned as shapes, and the result still validates
        vec = work / "vector.md"
        vec.write_text('## Title {layout=title}\nsub\n\n## Ops {layout=two_column icons="gear,cloud"}\n### Run\n- a\n### Host\n- b\n')
        run("fill_layout.py", vec, "-o", work / "vector.pptx")
        run("brand_qa.py", work / "vector.pptx")
        shapes = Presentation(str(work / "vector.pptx")).slides[1].shapes
        assert {"gear settings", "cloud hosting"} <= {sh.name for sh in shapes}

    # Deliberately injected defects are caught.
    bad = work / "bad.pptx"
    slide = prs.slides[1]
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    run_ = box.text_frame.paragraphs[0].add_run()
    run_.text = "off brand"
    run_.font.color.rgb = RGBColor(0xFF, 0x00, 0xFF)
    run_.font.name = "Comic Sans MS"
    slide.shapes.add_picture(str(fx / "stray.png"), Inches(6), Inches(5), Inches(1), Inches(1))
    prs.save(bad)
    shutil.copy(str(deck) + ".build.json", str(bad) + ".build.json")
    report = run("brand_qa.py", bad, expect=1)
    for needle in ("color FF00FF", "font 'Comic Sans MS'", "is not from the brand catalog"):
        assert needle in report, f"QA missed {needle!r}:\n{report}"

    shutil.rmtree(work)
    print("pptx-brand e2e: OK")


if __name__ == "__main__":
    main()
