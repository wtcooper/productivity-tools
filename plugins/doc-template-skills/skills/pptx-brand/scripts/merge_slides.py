#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5"]
# ///
"""Move pptxgenjs fallback slides onto a template layout inside the template-built deck.

Usage:
    merge_slides.py BASE.pptx FALLBACK.pptx -o OUT.pptx [--template NAME]
                    [--layout title_only] [--after N]

Each fallback slide becomes a new slide on the chosen template layout, so it picks up the
master's background, logo and footer instead of carrying a stray pptxgenjs master. Shapes,
pictures and native charts (with their embedded workbooks) are copied; the fallback slide's
own background is dropped. A text object created with `objectName: "brand-title"` is lifted
into the layout's title placeholder. Speaker notes are carried over.
"""

import argparse
import copy
import io
import re
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from _common import fail, load_json, ph_type, template_dir, write_json

R_ATTRS = ["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}" + a for a in ("embed", "link", "id", "pict")]
SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}


def adopt(part, package, seen):
    """Give a foreign part (and everything it relates to) a partname that is free in the destination package."""
    if id(part) in seen:
        return
    seen.add(id(part))
    template = re.sub(r"\d+(\.\w+)$", r"%d\1", str(part.partname))
    if "%d" in template:
        part.partname = package.next_partname(template)
    else:
        stem, _, ext = str(part.partname).rpartition(".")
        part.partname = package.next_partname(f"{stem}%d.{ext}")
    for rel in part.rels.values():
        if not rel.is_external:
            adopt(rel.target_part, package, seen)


def copy_slide(src, dest_prs, layout):
    new = dest_prs.slides.add_slide(layout)
    title_ph = next((p for p in new.placeholders if ph_type(p) in ("title", "ctrTitle")), None)
    for ph in list(new.placeholders):
        if title_ph is None or ph._element is not title_ph._element:  # proxies are rebuilt per access
            ph._element.getparent().remove(ph._element)

    rid_map, seen = {}, set()
    for rid, rel in src.part.rels.items():
        if rel.reltype in SKIP_RELS:
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.IMAGE:
            rid_map[rid] = new.part.get_or_add_image_part(io.BytesIO(rel.target_part.blob))[1]
        else:
            adopt(rel.target_part, dest_prs.part.package, seen)
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)

    title_text = None
    for shape in src.shapes:
        if shape.name == "brand-title" and shape.has_text_frame:
            title_text = shape.text_frame.text
            continue
        el = copy.deepcopy(shape._element)
        for node in el.iter():
            for attr in R_ATTRS:
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        new.shapes._spTree.append(el)
    for i, nv in enumerate(new.shapes._spTree.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr"), 1):
        nv.set("id", str(i))

    if title_ph is not None:
        if title_text:
            title_ph.text_frame.paragraphs[0].add_run().text = title_text
        else:
            title_ph._element.getparent().remove(title_ph._element)
    if src.has_notes_slide and src.notes_slide.notes_text_frame.text.strip():
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    return new


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("base")
    ap.add_argument("fallback")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--template")
    ap.add_argument("--layout", default="title_only", help="layout class, exact name, or #index (default title_only)")
    ap.add_argument("--after", type=int, help="insert after this slide number of BASE (default: append)")
    args = ap.parse_args()

    build_path = Path(args.base + ".build.json")
    build = load_json(build_path) if build_path.exists() else {}
    tdir = template_dir(args.template or build.get("template"))
    rows = load_json(tdir / "layouts.json")["layouts"]
    row = next((l for l in rows if args.layout.lower() in (l["class"], l["name"].lower(), f"#{l['index']}")), None) \
        or next((l for l in rows if l["class"] in ("title_only", "blank")), None)
    if row is None:
        fail(f"template has no {args.layout!r}, title_only or blank layout; pass --layout with an exact layout name")

    base, fallback = Presentation(args.base), Presentation(args.fallback)
    if (base.slide_width, base.slide_height) != (fallback.slide_width, fallback.slide_height):
        fail("slide sizes differ: call brand.applyLayout(pres) before adding slides in the pptxgenjs script")
    layout = next(l for m in base.slide_masters for l in m.slide_layouts if Path(l.part.partname).name == row["id"])

    start = len(base.slides)
    for src in fallback.slides:
        copy_slide(src, base, layout)
    if args.after is not None:
        ids = base.slides._sldIdLst
        moved = list(ids)[start:]
        for offset, el in enumerate(moved):
            ids.remove(el)
            ids.insert(args.after + offset, el)
    base.save(args.output)

    if build:
        # Images on fallback slides are deliberately not declared: brand_qa.py flags any that bypass the catalog.
        build["output"] = args.output
        write_json(args.output + ".build.json", build)
    print(f"wrote {args.output}: merged {len(fallback.slides)} slide(s) onto layout {row['name']!r}")


if __name__ == "__main__":
    main()
