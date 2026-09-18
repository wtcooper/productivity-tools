#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5"]
# ///
"""Markdown outline -> on-brand deck built from the template's own slide layouts.

Usage:
    fill_layout.py OUTLINE.md -o OUT.pptx [--template NAME]

Every slide is instantiated from a real slide layout and its placeholders are filled, so
the master's fonts, colors, logo and footer come along and the deck stays editable.
Outline syntax is documented in references/layout-mapping.md. A build report is printed
and saved as OUT.pptx.build.json (brand_qa.py reads it).
"""

import argparse
import copy
import io
import json
import math
import re
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn
from pptx.util import Inches

from _common import EMU_PER_IN, drop_all_slides, fail, load_json, ph_type, sha1, template_dir, write_json

FURNITURE = {"dt", "ftr", "sldNum", "hdr"}
TITLES = {"title", "ctrTitle"}
RASTER = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
ICON_IN, ICON_GAP_IN = 0.6, 0.15

FALLBACKS = {
    "three_column": ["two_column", "one_column"],
    "two_column": ["comparison", "one_column"],
    "comparison": ["two_column", "one_column"],
    "image_left": ["image_right", "two_column", "one_column"],
    "image_right": ["image_left", "two_column", "one_column"],
    "full_bleed_image": ["image_left", "image_right", "blank"],
    "chart": ["one_column", "title_only"],
    "table": ["one_column", "title_only"],
    "quote": ["section", "one_column"],
    "closing": ["title", "section"],
    "section": ["title", "title_only"],
    "title": ["section", "title_only"],
    "title_only": ["blank", "one_column"],
    "blank": ["title_only"],
    "one_column": ["two_column", "title_only"],
}

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


# ---------------------------------------------------------------- outline parsing

def parse_attrs(text):
    return {k: v.strip('"') for k, v in re.findall(r'(\w+)=("[^"]*"|\S+)', text)}


def parse_outline(md, base_dir):
    slides, slide, block, fence, in_notes = [], None, None, None, False

    def text_block(heading=None):
        nonlocal block
        block = {"kind": "text", "heading": heading, "paras": []}
        slide["blocks"].append(block)
        return block

    for raw in md.splitlines():
        line = raw.rstrip()
        if fence is not None:
            if line.strip() == "```":
                slide["blocks"].append({"kind": "chart", **json.loads("\n".join(fence))})
                fence, block = None, None
            else:
                fence.append(line)
            continue
        if line.startswith("## "):
            m = re.match(r"##\s+(.*?)\s*(\{(.*)\})?\s*$", line)
            slide = {"title": m.group(1), "attrs": parse_attrs(m.group(3) or ""), "blocks": [], "notes": []}
            slides.append(slide)
            block, in_notes = None, False
            continue
        if slide is None or not line.strip():
            continue
        stripped = line.strip()
        if in_notes:
            slide["notes"].append(stripped)
        elif re.match(r"notes?:", stripped, re.I):
            in_notes = True
            rest = stripped.split(":", 1)[1].strip()
            if rest:
                slide["notes"].append(rest)
        elif stripped.startswith("```chart"):
            fence = []
        elif stripped.startswith("### "):
            text_block(stripped[4:].strip())
        elif stripped == "---":
            text_block()
        elif m := re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped):
            ref = m.group(2)
            if not ref.startswith("asset:"):
                ref = str((base_dir / ref).resolve())
            slide["blocks"].append({"kind": "image", "alt": m.group(1), "ref": ref})
            block = None
        elif stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if block is None or block["kind"] != "table":
                block = {"kind": "table", "rows": []}
                slide["blocks"].append(block)
            block["rows"].append(cells)
        else:
            if block is None or block["kind"] != "text":
                text_block()
            m = re.match(r"(\s*)(?:[-*+]|\d+[.)])\s+(.*)", line)
            if m:
                block["paras"].append({"text": m.group(2), "level": len(m.group(1).replace("\t", "  ")) // 2, "bullet": True})
            else:
                block["paras"].append({"text": stripped.lstrip("> ").strip(), "level": 0, "bullet": False})
    if fence is not None:
        fail("unterminated ```chart block")
    return slides


# ---------------------------------------------------------------- layout resolution

def default_class(slide, position):
    kinds = [b["kind"] for b in slide["blocks"]]
    if position == 0:
        return "title"
    if not kinds:
        return "title_only"
    if "image" in kinds and "text" in kinds:
        return "image_right"
    if kinds == ["chart"]:
        return "chart"
    if kinds == ["table"]:
        return "table"
    return {1: "one_column", 2: "two_column"}.get(len(kinds), "three_column")


def resolve_layout(requested, layouts, n_blocks):
    """Exact name or '#index' wins; else the class, then its fallbacks. Returns (layout_row, substituted_from)."""
    for l in layouts:
        if requested.lower() == l["name"].lower() or requested == f"#{l['index']}":
            return l, None
    for cls in [requested] + FALLBACKS.get(requested, []) + ["one_column", "title_only", "blank"]:
        matches = [l for l in layouts if l["class"] == cls]
        if matches:
            def big_slots(l):
                return sum(1 for p in l["placeholders"] if p["type"] not in FURNITURE | TITLES | {"subTitle"})
            matches.sort(key=lambda l: (l.get("vertical", False), big_slots(l) < n_blocks, l["index"]))
            return matches[0], (None if cls == requested else requested)
    fail(f"no usable layout for {requested!r}")


# ---------------------------------------------------------------- assets

class Assets:
    def __init__(self, tdir):
        self.tdir = tdir
        path = tdir / "assets.json"
        self.rows = load_json(path)["assets"] if path.exists() else []

    def find(self, query):
        query = query.strip().lower().removeprefix("asset:")
        if query.lstrip("#").isdigit():
            return next((a for a in self.rows if a["index"] == int(query.lstrip("#"))), None)
        best, best_score = None, 0
        for a in self.rows:
            if a["format"] not in RASTER | {"shape"} and not a["preview"]:
                continue  # vector file that could not be rasterized at onboarding (no LibreOffice)
            score = 3 * (query in [t.lower() for t in a["tags"]]) + 2 * (query in a["label"].lower().split()) \
                + (query in a["label"].lower())
            if score > best_score:
                best, best_score = a, score
        return best

    def image_bytes(self, asset, tint_hex=None):
        """Bytes python-pptx can embed. Vector files use their PNG preview; recolorable art can be tinted."""
        path = self.tdir / asset["file"]
        if asset["format"] not in RASTER:
            path = self.tdir / asset["preview"]
        data = path.read_bytes()
        if tint_hex and asset["recolorable"]:
            img = Image.open(io.BytesIO(data)).convert("RGBA")
            solid = Image.new("RGBA", img.size, tuple(int(tint_hex[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
            solid.putalpha(img.getchannel("A"))
            buf = io.BytesIO()
            solid.save(buf, "PNG")
            data = buf.getvalue()
        return data


# ---------------------------------------------------------------- slide filling

def add_runs(paragraph, text, bold=False):
    """**bold** spans become bold runs; formatting otherwise inherits from the layout."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run = paragraph.add_run()
            run.text = part
            if bold or i % 2:
                run.font.bold = True


def write_text(ph, heading, paras, plain_unbulleted):
    tf = ph.text_frame
    first = True
    items = ([{"text": heading, "level": 0, "bullet": False, "bold": True}] if heading else []) + paras
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = min(item["level"], 8)
        add_runs(p, item["text"], bold=item.get("bold", False))
        if plain_unbulleted and not item["bullet"]:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", "0")
            ppr.set("indent", "0")
            etree.SubElement(ppr, qn("a:buNone"))


def box_of(ph):
    return ph.left, ph.top, ph.width, ph.height


def remove(shape):
    shape._element.getparent().remove(shape._element)


def fit_picture(slide, data, box):
    left, top, width, height = box
    w, h = Image.open(io.BytesIO(data)).size
    scale = min(width / w, height / h)
    pw, ph_ = int(w * scale), int(h * scale)
    return slide.shapes.add_picture(io.BytesIO(data), left + (width - pw) // 2, top + (height - ph_) // 2, pw, ph_)


def add_chart(slide, spec, box, brand):
    data = CategoryChartData()
    data.categories = spec["categories"]
    for s in spec["series"]:
        data.add_series(s["name"], s["values"])
    kind = spec.get("type", "column")
    if kind not in CHART_TYPES:
        fail(f"chart type {kind!r} not supported; use one of {sorted(CHART_TYPES)}")
    chart = slide.shapes.add_chart(CHART_TYPES[kind], *box, data).chart
    chart.has_legend = len(spec["series"]) > 1 or kind in ("pie", "doughnut")
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if spec.get("title"):
        chart.has_title = True
        chart.chart_title.text_frame.text = spec["title"]
    else:
        chart.has_title = False
    if kind in ("pie", "doughnut"):
        chart.plots[0].has_data_labels = True
        return
    # Series colors are theme references, so they follow the template and never hard-code a hex.
    for series, slot in zip(chart.plots[0].series, brand["color_roles"]["chart_series_order"]):
        theme = getattr(MSO_THEME_COLOR, slot.upper().replace("ACCENT", "ACCENT_"))
        fills = [series.format.fill]
        if kind == "line":
            series.format.line.color.theme_color = theme
            fills = [series.marker.format.fill]
        for fill in fills:
            fill.solid()
            fill.fore_color.theme_color = theme


def add_table(slide, rows, box):
    cols = max(len(r) for r in rows)
    left, top, width, height = box
    row_h = min(height // len(rows), Inches(0.5))
    table = slide.shapes.add_table(len(rows), cols, left, top, width, row_h * len(rows)).table
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            para = table.cell(r, c).text_frame.paragraphs[0]
            add_runs(para, text)


def clone_shape(slide, xml_bytes, left, top, size):
    """Vector icon harvested from an icon deck: clone its XML and scale it into place."""
    el = etree.fromstring(xml_bytes)
    xfrm = el.find(".//" + qn("a:xfrm"))
    ext = xfrm.find(qn("a:ext"))
    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
    scale = size / max(cx, cy)
    xfrm.find(qn("a:off")).set("x", str(int(left)))
    xfrm.find(qn("a:off")).set("y", str(int(top)))
    ext.set("cx", str(int(cx * scale)))
    ext.set("cy", str(int(cy * scale)))
    slide.shapes._spTree.append(el)
    for nv in el.iter(qn("p:cNvPr")):
        nv.set("id", str(slide.shapes._next_shape_id))


class Builder:
    def __init__(self, tdir):
        self.tdir = tdir
        self.brand = load_json(tdir / "brand.json")
        self.layout_rows = load_json(tdir / "layouts.json")["layouts"]
        self.assets = Assets(tdir)
        self.prs = Presentation(str(tdir / "source.pptx"))
        self.by_id = {Path(l.part.partname).name: l for m in self.prs.slide_masters for l in m.slide_layouts}
        drop_all_slides(self.prs)
        self.slide_h = self.prs.slide_height
        self.report, self.derived = [], []

    def slots(self, slide):
        """(big content slots with their paired heading slot, leftover small text slots), in reading order."""
        content = [p for p in slide.placeholders if ph_type(p) not in FURNITURE | TITLES | {"subTitle"} and p.width]
        big = sorted([p for p in content if p.height >= 0.25 * self.slide_h], key=lambda p: (p.left, p.top))
        small = [p for p in content if p not in big]
        pairs = []
        for slot in big:
            head = next((s for s in small if ph_type(s) == "body" and abs(s.left - slot.left) < 0.3 * EMU_PER_IN
                         and s.top < slot.top), None)
            if head is not None:
                small.remove(head)
            pairs.append([slot, head])
        return pairs, sorted(small, key=lambda p: (p.top, p.left))

    def split_for_capacity(self, spec, row):
        """A single text block longer than its placeholder holds continues on a duplicated layout."""
        texts = [b for b in spec["blocks"] if b["kind"] == "text"]
        caps = [p for p in row["placeholders"] if "capacity_lines" in p and p["box_in"][3] >= 0.25 * self.brand["slide_size_in"]["h"]]
        if len(spec["blocks"]) != 1 or len(texts) != 1 or len(caps) != 1:
            return [spec]
        cap, width = caps[0]["capacity_lines"], caps[0]["chars_per_line"]
        groups = []
        for para in texts[0]["paras"]:
            if para["level"] == 0 or not groups:
                groups.append([])
            groups[-1].append(para)
        chunks, current, used = [], [], 1 if texts[0]["heading"] else 0
        for group in groups:
            need = sum(max(1, math.ceil(len(p["text"]) / max(1, width - 4 * p["level"]))) for p in group)
            if current and used + need > cap:
                chunks.append(current)
                current, used = [], 0
            current += group
            used += need
        chunks.append(current)
        if len(chunks) == 1:
            return [spec]
        out = []
        for i, paras in enumerate(chunks):
            part = copy.deepcopy(spec)
            part["blocks"][0]["paras"] = paras
            if i:
                part["title"] = f"{spec['title']} (cont.)"
                part["attrs"].pop("icons", None)
            out.append(part)
        return out

    def place_icons(self, slide, spec, pairs, free_pics, entry):
        queries = [q for q in spec["attrs"].get("icons", "").split(",") if q.strip()]
        tint = spec["attrs"].get("icon_color")
        tint_hex = self.brand["colors"].get(self.brand["color_roles"].get(tint, tint)) if tint else None
        if tint and not tint_hex:
            entry["warnings"].append(f"icon_color {tint!r} is not a color role or theme slot")
        columns = [pair for pair in pairs if pair[0] is not None]
        for i, query in enumerate(queries):
            asset = self.assets.find(query)
            if asset is None:
                entry["unresolved_icons"].append(query.strip())
                continue
            entry["assets"].append(f"#{asset['index']} {asset['label']}")
            if asset["format"] != "shape":
                data = self.assets.image_bytes(asset, tint_hex)
                self.derived.append(sha1(data))
            if free_pics and asset["format"] != "shape":
                free_pics.pop(0).insert_picture(io.BytesIO(data))
                continue
            if i >= len(columns):
                entry["warnings"].append(f"no free region for icon {query.strip()!r}")
                continue
            slot, head = columns[i]
            anchor = head if head is not None else slot
            left, top = anchor.left, anchor.top
            shift = Inches(ICON_IN + ICON_GAP_IN)
            for shape in filter(None, (head, slot)):
                # Set all four: a placeholder that inherits its geometry has none of its own to adjust.
                l, t, w, h = box_of(shape)
                shape.left, shape.top, shape.width = l, t + shift, w
                shape.height = h - shift if shape is slot else h
            if asset["format"] == "shape":
                clone_shape(slide, (self.tdir / asset["file"]).read_bytes(), left, top, Inches(ICON_IN))
            else:
                fit_picture(slide, data, (left, top, Inches(ICON_IN), Inches(ICON_IN)))

    def build_slide(self, spec, position):
        requested = spec["attrs"].get("layout") or default_class(spec, position)
        row, substituted = resolve_layout(requested, self.layout_rows, len(spec["blocks"]))
        parts = self.split_for_capacity(spec, row)
        for part in parts:
            self.fill(part, row, requested, substituted)

    def fill(self, spec, row, requested, substituted):
        slide = self.prs.slides.add_slide(self.by_id[row["id"]])
        entry = {"slide": len(self.prs.slides), "title": spec["title"], "layout": row["name"], "class": row["class"],
                 "requested": requested, "substituted": bool(substituted), "assets": [], "unresolved_icons": [],
                 "warnings": []}
        self.report.append(entry)
        if substituted:
            entry["warnings"].append(f"template has no {requested!r} layout; used {row['class']!r} ({row['name']})")

        used = set()
        title = next((p for p in slide.placeholders if ph_type(p) in TITLES), None)
        if title is not None:
            add_runs(title.text_frame.paragraphs[0], spec["title"])
            used.add(title.shape_id)
        elif spec["title"]:
            entry["warnings"].append("layout has no title placeholder; title omitted")

        blocks = list(spec["blocks"])
        subtitle = next((p for p in slide.placeholders if ph_type(p) == "subTitle"), None)
        if subtitle is not None and blocks and blocks[0]["kind"] == "text":
            b = blocks.pop(0)
            write_text(subtitle, b["heading"], b["paras"], plain_unbulleted=False)
            used.add(subtitle.shape_id)

        pairs, small = self.slots(slide)
        pic_slots = [p for p, _ in pairs if ph_type(p) == "pic"]
        columns = [pair for pair in pairs if ph_type(pair[0]) != "pic"]
        queue = columns + [[s, None] for s in small]
        last_text = None
        for b in blocks:
            if b["kind"] == "image":
                data = self.image_data(b, entry)
                if data is None:
                    continue
                if pic_slots:
                    ph = pic_slots.pop(0)
                    used.add(ph.shape_id)
                    ph.insert_picture(io.BytesIO(data))
                elif queue:
                    slot, _ = queue.pop(0)
                    fit_picture(slide, data, box_of(slot))
                else:
                    entry["warnings"].append(f"no placeholder left for image {b['ref']!r}")
                continue
            if not queue:
                if b["kind"] == "text" and last_text is not None:
                    extra = ([{"text": b["heading"], "level": 0, "bullet": False, "bold": True}] if b["heading"] else []) + b["paras"]
                    for item in extra:
                        p = last_text.text_frame.add_paragraph()
                        p.level = item["level"]
                        add_runs(p, item["text"], bold=item.get("bold", False))
                    entry["warnings"].append("more text blocks than placeholders; extra block appended to the last one")
                else:
                    entry["warnings"].append(f"no placeholder left for {b['kind']} block")
                continue
            slot, head = queue.pop(0)
            if b["kind"] == "text":
                heading = b["heading"]
                if head is not None and heading:
                    add_runs(head.text_frame.paragraphs[0], heading)
                    used.add(head.shape_id)
                    heading = None
                write_text(slot, heading, b["paras"], plain_unbulleted=slot.height >= 0.25 * self.slide_h)
                used.add(slot.shape_id)
                last_text = slot
                cap = next((p for p in row["placeholders"] if p["idx"] == slot.placeholder_format.idx), {})
                if "capacity_lines" in cap:
                    need = sum(max(1, math.ceil(len(p["text"]) / cap["chars_per_line"])) for p in b["paras"]) + bool(heading)
                    if need > cap["capacity_lines"]:
                        entry["warnings"].append(f"block {b['heading'] or b['paras'][0]['text'][:30]!r} needs ~{need} lines; "
                                                 f"placeholder holds ~{cap['capacity_lines']}: trim or split the slide")
            elif b["kind"] == "chart":
                add_chart(slide, b, box_of(slot), self.brand)
            else:
                add_table(slide, b["rows"], box_of(slot))

        free_pics = [p for p in pic_slots]
        self.place_icons(slide, spec, [pair for pair in columns if pair[0].shape_id in used], free_pics, entry)
        used.update(p.shape_id for p in pic_slots if p not in free_pics)

        for ph in list(slide.placeholders):
            if ph.shape_id not in used and ph_type(ph) not in FURNITURE:
                remove(ph)
        notes = list(spec["notes"])
        if entry["assets"]:
            notes.append("Brand catalog assets: " + "; ".join(entry["assets"]))
        if notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)

    def image_data(self, block, entry):
        ref = block["ref"]
        if ref.startswith("asset:"):
            asset = self.assets.find(ref)
            if asset is None or asset["format"] == "shape":
                entry["warnings"].append(f"{ref} is not an image in the catalog")
                return None
            entry["assets"].append(f"#{asset['index']} {asset['label']}")
            data = self.assets.image_bytes(asset)
        else:
            if not Path(ref).is_file():
                entry["warnings"].append(f"image not found: {ref}")
                return None
            data = Path(ref).read_bytes()
        self.derived.append(sha1(data))
        return data


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("outline")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--template")
    args = ap.parse_args()

    tdir = template_dir(args.template)
    outline = Path(args.outline)
    slides = parse_outline(outline.read_text(), outline.parent)
    if not slides:
        fail("outline has no '## ' slide headings")
    builder = Builder(tdir)
    for position, spec in enumerate(slides):
        builder.build_slide(spec, position)
    builder.prs.save(args.output)

    report = {"template": tdir.name, "output": args.output, "slides": builder.report,
              "embedded_image_sha1": sorted(set(builder.derived))}
    write_json(args.output + ".build.json", report)
    warnings = sum(len(s["warnings"]) + len(s["unresolved_icons"]) for s in builder.report)
    for s in builder.report:
        flags = "; ".join(s["warnings"] + [f"unresolved icon {q!r}" for q in s["unresolved_icons"]])
        print(f"{s['slide']:>3}  {s['layout']:<28} {s['title'][:40]:<40} {flags}")
    print(f"wrote {args.output}: {len(builder.report)} slides, {warnings} warning(s). Next: brand_qa.py {args.output}")


if __name__ == "__main__":
    main()
