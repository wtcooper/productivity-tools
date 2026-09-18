#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5", "defusedxml"]
# ///
"""Brand, asset, layout, content and file QA for a deck built from a registered template.

Usage:
    brand_qa.py DECK.pptx [--template NAME] [--allow-image FILE]... [--json]

Checks:
    brand    every hard-coded color and font on slides/charts is in brand.json
    asset    every embedded image is a catalog asset, one the outline supplied, or an --allow-image file
    layout   every slide sits on a layout profiled from the template (no stray masters)
    content  no leftover placeholder text or empty placeholders
    file     the official pptx skill's validate.py --original, when that skill is installed

Exit status is 1 when any check fails.
"""

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from _common import find_upstream, load_json, sha1, template_dir

LEFTOVER = re.compile(r"\bx{3,}\b|lorem|ipsum|\bTODO\b|\[insert|click to (add|edit)", re.I)


def check_brand(prs, brand):
    overrides = brand.get("overrides") or {}
    colors = ({v.upper() for v in brand["colors"].values()} | {c["hex"].upper() for c in brand["custom_colors"]}
              | {c.upper() for c in overrides.get("extra_colors", [])} | {"FFFFFF", "000000"})
    fonts = ({f.lower() for t in brand["themes"] for f in t["fonts"].values()}
             | {brand["font_safety"]["fallback_body"].lower()} | {f.lower() for f in overrides.get("extra_fonts", [])})
    findings = []
    for n, slide in enumerate(prs.slides, 1):
        parts = [slide.part] + [r.target_part for r in slide.part.rels.values() if r.reltype == RT.CHART]
        for part in parts:
            xml = part.blob.decode("utf-8", "ignore")
            where = f"slide {n}" + ("" if part is slide.part else f" ({Path(part.partname).name})")
            for hex_ in sorted(set(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml))):
                if hex_.upper() not in colors:
                    findings.append(f"{where}: color {hex_.upper()} is not a brand color")
            for face in sorted(set(re.findall(r'<a:(?:latin|ea|cs) typeface="([^"]+)"', xml))):
                if not face.startswith("+") and face.lower() not in fonts:
                    findings.append(f"{where}: font {face!r} is not a brand font")
    return findings


def check_assets(prs, tdir, declared):
    allowed = set(declared)
    for path in (tdir / "assets").rglob("*"):
        if path.is_file():
            allowed.add(sha1(path.read_bytes()))
    findings = []
    for n, slide in enumerate(prs.slides, 1):
        for rel in slide.part.rels.values():
            if rel.reltype == RT.IMAGE and not rel.is_external and sha1(rel.target_part.blob) not in allowed:
                findings.append(f"slide {n}: image {Path(rel.target_part.partname).name} is not from the brand catalog "
                                "(use a catalog asset when a tagged match exists)")
    return findings


def check_layouts(prs, layouts):
    known = {(l["id"], l["name"]) for l in layouts}
    findings = []
    for n, slide in enumerate(prs.slides, 1):
        layout = slide.slide_layout
        if (Path(layout.part.partname).name, layout.name) not in known:
            findings.append(f"slide {n}: layout {layout.name!r} ({Path(layout.part.partname).name}) is not a profiled "
                            "template layout; merge the slide with merge_slides.py")
    return findings


def check_content(prs):
    findings = []
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            if LEFTOVER.search(text):
                findings.append(f"slide {n}: leftover placeholder text in {shape.name!r}: {text[:50]!r}")
            elif shape.is_placeholder and not text.strip():
                findings.append(f"slide {n}: empty placeholder {shape.name!r}; fill it or delete it")
    return findings


def check_file(deck, tdir):
    upstream = find_upstream()
    if upstream is None:
        return None, "skipped: official pptx skill not found (set PPTX_SKILL_DIR or install document-skills)"
    proc = subprocess.run([sys.executable, str(upstream / "scripts/office/validate.py"), str(deck),
                           "--original", str(tdir / "source.pptx")], capture_output=True, text=True,
                          cwd=upstream / "scripts/office")
    output = (proc.stdout + proc.stderr).strip()
    return ([] if proc.returncode == 0 else [output[-2000:]]), output.splitlines()[-1] if output else ""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--template")
    ap.add_argument("--allow-image", action="append", default=[], help="a non-catalog image that is legitimately on a slide")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    deck = Path(args.deck).resolve()
    report_path = Path(str(deck) + ".build.json")
    build = load_json(report_path) if report_path.exists() else {}
    tdir = template_dir(args.template or build.get("template"))
    prs = Presentation(str(deck))

    results = {
        "brand": check_brand(prs, load_json(tdir / "brand.json")),
        "asset": check_assets(prs, tdir, build.get("embedded_image_sha1", [])
                              + [sha1(Path(f).read_bytes()) for f in args.allow_image]),
        "layout": check_layouts(prs, load_json(tdir / "layouts.json")["layouts"]),
        "content": check_content(prs),
    }
    file_findings, file_note = check_file(deck, tdir)
    if file_findings is not None:
        results["file"] = file_findings
    failed = any(results.values())

    if args.json:
        print(json.dumps({"deck": str(deck), "template": tdir.name, "passed": not failed, "findings": results,
                          "file_check": file_note}, indent=2))
    else:
        for name, findings in results.items():
            print(f"{name:<8} {'FAIL' if findings else 'ok'}")
            for f in findings:
                print(f"         - {f}")
        if file_findings is None:
            print(f"file     {file_note}")
        print("QA FAILED" if failed else "QA passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
