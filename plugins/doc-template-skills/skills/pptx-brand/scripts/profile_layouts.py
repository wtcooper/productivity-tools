#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5"]
# ///
"""Slide masters/layouts -> layouts.json + labeled layout thumbnails.

Usage:
    profile_layouts.py <template-dir> [--no-render]

Records every layout's placeholders (type, idx, inch geometry, capacity), the
decorative regions that are not available for content, and a coarse layout class.
Rows marked `"edited": true` in an existing layouts.json keep their class/notes/use_for.
"""

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from _common import can_render, contact_sheets, drop_all_slides, emu_to_in, load_json, ph_type, render_slides, write_json

ALL_CLASSES = ["title", "section", "one_column", "two_column", "three_column", "image_left", "image_right",
               "full_bleed_image", "chart", "table", "comparison", "quote", "closing", "title_only", "blank"]

USE_FOR = {
    "title": "the opening slide: deck title and subtitle",
    "section": "a divider that introduces a new section",
    "one_column": "one idea with supporting bullets, or a single chart/table",
    "two_column": "two parallel ideas, or text beside a chart",
    "three_column": "three parallel points, pillars, or options",
    "image_left": "a picture on the left with explanatory text",
    "image_right": "explanatory text with a picture on the right",
    "full_bleed_image": "a single full-slide photo, optionally with a short title",
    "chart": "one native chart",
    "table": "one table",
    "comparison": "two labeled columns: before/after, option A vs option B",
    "quote": "a pull quote or customer statement",
    "closing": "the final slide: thanks, contacts, next steps",
    "title_only": "a title above free space for a bespoke visual (fallback slides land here)",
    "blank": "free canvas with only the master's decoration",
}

NAME_HINTS = [(r"clos|thank|final|end slide", "closing"), (r"quote", "quote"),
              (r"section|divider|chapter", "section"), (r"compar|versus|\bvs\b", "comparison")]

FURNITURE = {"dt", "ftr", "sldNum", "hdr"}
TITLES = {"title", "ctrTitle"}
DEFAULT_PT = {"title": 40, "ctrTitle": 44, "subTitle": 24}


def _size_from(el, path):
    found = el.xpath(path) if el is not None else []
    return int(found[0]) / 100 if found else None


def default_font_pt(shape, kind, master):
    """Default level-1 size: the placeholder, then the master's matching placeholder, then master txStyles."""
    size = _size_from(shape._element, ".//a:lstStyle/a:lvl1pPr/a:defRPr/@sz")
    if size:
        return size
    want = "title" if kind in TITLES else "body"
    for ph in master.placeholders:
        if ph_type(ph) == want:
            size = _size_from(ph._element, ".//a:lstStyle/a:lvl1pPr/a:defRPr/@sz")
            if size:
                return size
    style = "titleStyle" if kind in TITLES else "bodyStyle"
    size = _size_from(master.element, f"./p:txStyles/p:{style}/a:lvl1pPr/a:defRPr/@sz")
    return size or DEFAULT_PT.get(kind, 18)


def capacity(box_in, pt, kind):
    """Rough fit at the default size: average glyph ~0.45em wide, lines 1.2em tall, small insets."""
    w, h = box_in[2] - 0.2, box_in[3] - 0.1
    chars_per_line = max(1, int(w * 72 / (pt * 0.45)))
    lines = max(1, int(h * 72 / (pt * 1.2) * (1.0 if kind in TITLES else 0.85)))
    if kind in TITLES or kind == "subTitle":
        return {"capacity_chars": chars_per_line * min(lines, 2)}
    return {"capacity_lines": lines, "chars_per_line": chars_per_line}


def classify(name, phs, slide_w, slide_h):
    lowered = name.lower()
    for pattern, cls in NAME_HINTS:
        if re.search(pattern, lowered):
            return cls
    kinds = [p["type"] for p in phs]
    content = [p for p in phs if p["type"] not in FURNITURE | TITLES | {"subTitle"} and p.get("box_in")]
    big = [p for p in content if p["box_in"][3] >= 0.25 * slide_h]
    small_text = [p for p in content if p not in big and p["type"] == "body"]

    if "ctrTitle" in kinds or ("subTitle" in kinds and not big):
        return "title"
    if not big:
        if small_text:
            return "section"
        return "title_only" if "title" in kinds else "blank"
    for p in big:
        if p["type"] == "pic" and p["box_in"][2] * p["box_in"][3] >= 0.8 * slide_w * slide_h:
            return "full_bleed_image"
    columns = sorted({round(p["box_in"][0], 1) for p in big})
    if len(columns) >= 3:
        return "three_column"
    if len(columns) == 2:
        if len(small_text) >= 2:
            return "comparison"
        pics = [p for p in big if p["type"] == "pic"]
        if len(pics) == 1:
            return "image_left" if round(pics[0]["box_in"][0], 1) == columns[0] else "image_right"
        return "two_column"
    only = big[0]
    if only["type"] == "pic":
        return "image_left" if only["box_in"][0] + only["box_in"][2] / 2 < slide_w / 2 else "image_right"
    return {"chart": "chart", "tbl": "table"}.get(only["type"], "one_column")


def reserved_regions(shapes, slide_w, slide_h):
    """Non-placeholder shapes are decoration. Near-full-slide ones are background art, not a reserved box."""
    regions, background = [], False
    for shape in shapes:
        if shape.is_placeholder or shape.width is None or not shape.width or not shape.height:
            continue
        box = [emu_to_in(v) for v in (shape.left, shape.top, shape.width, shape.height)]
        if box[2] * box[3] >= 0.6 * slide_w * slide_h:
            background = True
        else:
            regions.append(box)
    return regions, background


def profile(tdir):
    prs = Presentation(str(Path(tdir) / "source.pptx"))
    slide_w, slide_h = emu_to_in(prs.slide_width), emu_to_in(prs.slide_height)
    layouts = []
    for master in prs.slide_masters:
        master_regions, master_bg = reserved_regions(master.shapes, slide_w, slide_h)
        for layout in master.slide_layouts:
            phs = []
            for ph in layout.placeholders:
                kind = ph_type(ph)
                row = {"type": kind, "idx": ph.placeholder_format.idx, "name": ph.name,
                       "inherits_geometry": not ph._element.xpath("./p:spPr/a:xfrm")}
                if ph.width:
                    row["box_in"] = [emu_to_in(v) for v in (ph.left, ph.top, ph.width, ph.height)]
                    if kind not in FURNITURE | {"pic", "chart", "tbl", "media", "clipArt", "dgm"}:
                        pt = default_font_pt(ph, kind, master)
                        row["default_pt"] = pt
                        row.update(capacity(row["box_in"], pt, kind))
                phs.append(row)
            regions, bg = reserved_regions(layout.shapes, slide_w, slide_h)
            cls = classify(layout.name, phs, slide_w, slide_h)
            layouts.append({
                "id": Path(layout.part.partname).name,
                "index": len(layouts) + 1,
                "name": layout.name,
                "class": cls,
                "master": Path(master.part.partname).name,
                "use_for": USE_FOR[cls],
                "placeholders": phs,
                "reserved_regions_in": regions + master_regions,
                "has_background_art": bg or master_bg,
                "vertical": bool(layout.element.xpath(".//a:bodyPr[@vert and @vert!='horz']")),
                "notes": "",
            })
    return prs, layouts


def render_layouts(tdir, layouts):
    """One scratch slide per layout, each placeholder labeled with its own type and idx."""
    tdir = Path(tdir)
    prs = Presentation(str(tdir / "source.pptx"))
    drop_all_slides(prs)
    all_layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
    for layout in all_layouts:
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            label = f"{ph_type(ph).upper()} idx={ph.placeholder_format.idx}"
            if ph.has_text_frame and ph_type(ph) not in ("pic", "chart", "tbl"):
                ph.text_frame.text = label
            elif ph.width:
                box = slide.shapes.add_textbox(ph.left, ph.top, ph.width, ph.height)
                box.text_frame.text = label
                run = box.text_frame.paragraphs[0].runs[0]
                box.line.color.rgb = run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
                run.font.size = Pt(14)
    thumbs = tdir / "thumbs"
    thumbs.mkdir(exist_ok=True)
    scratch = thumbs / "layouts-scratch.pptx"
    prs.save(str(scratch))
    pngs = render_slides(scratch, thumbs, prefix="layout")
    scratch.unlink()
    for old in thumbs.glob("layouts-*.jpg"):
        old.unlink()
    items = [(f"{l['index']}. {l['name']} [{l['class']}]", png) for l, png in zip(layouts, pngs)]
    sheets = contact_sheets(items, thumbs / "layouts", cols=3, cell_w=420, per_sheet=12)
    for i, (l, png) in enumerate(zip(layouts, pngs)):
        l["png"] = f"thumbs/{png.name}"
        l["thumb"] = f"thumbs/{sheets[i // 12].name}#{l['index']}"


def build_layouts(tdir, render=True):
    tdir = Path(tdir)
    prs, layouts = profile(tdir)
    out = tdir / "layouts.json"
    if out.exists():
        edited = {l["id"]: l for l in load_json(out)["layouts"] if l.get("edited")}
        for l in layouts:
            if l["id"] in edited:
                l.update({k: edited[l["id"]][k] for k in ("class", "use_for", "notes", "edited")})
    if render and can_render():
        render_layouts(tdir, layouts)
    present = {l["class"] for l in layouts}
    coverage = {"missing_classes": [c for c in ALL_CLASSES if c not in present]}
    content_layouts = [l for l in layouts if l["class"] not in ("title", "section", "closing", "title_only", "blank")]
    if len(content_layouts) < 2 and len(prs.slides) >= 5:
        coverage["warning"] = ("Few content layouts but many example slides: this file is likely an example deck whose "
                               "decoration lives on slides, not layouts. Promote representative slides to layouts in "
                               "PowerPoint's Slide Master view and onboard again, or output will look flat.")
    write_json(out, {"layouts": layouts, "coverage": coverage})
    return layouts, coverage


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template_dir")
    ap.add_argument("--no-render", action="store_true")
    args = ap.parse_args()
    layouts, coverage = build_layouts(args.template_dir, render=not args.no_render)
    print(f"wrote layouts.json: {len(layouts)} layouts; missing classes: {', '.join(coverage['missing_classes']) or 'none'}")
    if not args.no_render and not can_render():
        print("warning: LibreOffice (soffice) or pdftoppm not found; layout thumbnails skipped")
    if "warning" in coverage:
        print("warning: " + coverage["warning"])


if __name__ == "__main__":
    main()
