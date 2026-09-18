#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2", "Pillow>=10", "lxml>=5"]
# ///
"""Template media + icon/graphic folders + icon-library decks -> assets.json + contact sheets.

Usage:
    harvest_assets.py <template-dir> [--icons DIR]... [--graphics DIR]... [--icon-deck FILE.pptx]...

Assets are copied into <template-dir>/assets/ so the template folder is self-contained.
Labels, tags and kind are keyed by content hash: re-running never discards a label
that was written for an unchanged file.
"""

import argparse
import copy
import re
import shutil
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from _common import (can_render, contact_sheets, emu_to_in, load_json, render_slides, sha1, to_png,
                     write_json)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".svg", ".emf", ".wmf"}
VECTOR_EXTS = {".svg", ".emf", ".wmf"}


def guess_label(stem):
    words = [w for w in re.split(r"[\W_]+|(?<=[a-z])(?=[A-Z])", stem) if w and not re.fullmatch(r"s?\d+", w)]
    drop = {"icon", "ic", "img", "image", "picture", "graphic", "svg", "png"}
    words = [w.lower() for w in words if w.lower() not in drop]
    return " ".join(words), words


def dominant_colors(png_path, limit=3):
    img = Image.open(png_path).convert("RGBA")
    img.thumbnail((64, 64))
    raw = img.tobytes()
    pixels = [tuple(raw[i:i + 3]) for i in range(0, len(raw), 4) if raw[i + 3] > 128]
    if not pixels:
        return []
    counts = {}
    for r, g, b in pixels:
        key = (r // 32, g // 32, b // 32)
        counts.setdefault(key, []).append((r, g, b))
    top = sorted(counts.values(), key=len, reverse=True)[:limit]
    keep = [grp for grp in top if len(grp) >= 0.08 * len(pixels)]
    return ["%02X%02X%02X" % tuple(sum(c[i] for c in grp) // len(grp) for i in range(3)) for grp in keep]


def is_recolorable(path, png_path):
    """Single-color art on transparency can be tinted to a brand token without losing anything."""
    if path.suffix.lower() == ".svg":
        fills = set(re.findall(r"(?:fill|stroke)\s*[:=]\s*[\"']?(#[0-9a-fA-F]{3,6}|currentColor)", path.read_text(errors="ignore")))
        return len(fills) <= 1
    img = Image.open(png_path).convert("RGBA")
    has_alpha = img.getextrema()[3][0] < 255
    return has_alpha and len(dominant_colors(png_path)) == 1


def guess_kind(path, source, px):
    if source in ("icons", "icon-deck"):
        return "icon"
    name = path.stem.lower()
    if "logo" in name:
        return "logo"
    if path.suffix.lower() in VECTOR_EXTS or max(px or (0, 0)) <= 512:
        return "icon"
    if path.suffix.lower() in (".jpg", ".jpeg"):
        return "photo"
    return "pattern" if re.search(r"pattern|texture|bg|background", name) else "diagram"


class Harvest:
    def __init__(self, tdir):
        self.tdir = Path(tdir)
        self.root = self.tdir / "assets"
        self.previews = self.root / "previews"
        self.previews.mkdir(parents=True, exist_ok=True)
        old = load_json(self.tdir / "assets.json")["assets"] if (self.tdir / "assets.json").exists() else []
        self.cache = {a["sha1"]: a for a in old}
        self.assets = {}

    def add(self, data, rel_path, source, used_by=None, shape_xml=None, size_in=None, preview=None):
        digest = sha1(data)
        if digest in self.assets:
            self.assets[digest]["used_by"] = sorted(set(self.assets[digest]["used_by"]) | set(used_by or []))
            return
        dest = self.root / rel_path
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        png = self.previews / f"{digest[:12]}.png"
        if preview is not None:
            preview.save(png)
        elif not png.exists() and not to_png(dest, png):
            png = None
        px = Image.open(png).size if png else None
        if dest.suffix.lower() not in VECTOR_EXTS and shape_xml is None and png:
            px = Image.open(dest).size
        label, tags = guess_label(dest.stem)
        row = {
            "sha1": digest,
            "file": str(dest.relative_to(self.tdir)),
            "preview": str(png.relative_to(self.tdir)) if png else None,
            "source": source,
            "kind": guess_kind(dest, source, px),
            "format": "shape" if shape_xml else dest.suffix.lower().lstrip("."),
            "label": label, "tags": tags, "labeled": False,
            "native_colors": dominant_colors(png) if png else [],
            "recolorable": bool(png) and shape_xml is None and is_recolorable(dest, png),
            "px": list(px) if px else None,
            "size_in": size_in,
            "used_by": sorted(used_by or []),
        }
        cached = self.cache.get(digest)
        if cached and cached.get("labeled"):
            row.update({k: cached[k] for k in ("label", "tags", "kind", "labeled")})
        self.assets[digest] = row

    def template_media(self):
        """Every image the template ships, with the layouts/masters/slides that reference it."""
        prs = Presentation(str(self.tdir / "source.pptx"))
        holders = list(prs.slide_masters) + [l for m in prs.slide_masters for l in m.slide_layouts] + list(prs.slides)
        for holder in holders:
            for rel in holder.part.rels.values():
                if rel.reltype == RT.IMAGE and not rel.is_external:
                    part = rel.target_part
                    self.add(part.blob, f"media/{Path(part.partname).name}", "template",
                             used_by=[Path(holder.part.partname).name])

    def folder(self, directory, source):
        base = Path(directory)
        for path in sorted(p for p in base.rglob("*") if p.suffix.lower() in IMAGE_EXTS):
            self.add(path.read_bytes(), f"{source}/{path.relative_to(base)}", source)

    def icon_deck(self, deck_path):
        """Pictures are exported as files. Pure-vector shapes/groups keep their XML so they can be cloned, not rasterized."""
        deck_path = Path(deck_path)
        prs = Presentation(str(deck_path))
        renders = render_slides(deck_path, self.previews / "_deck", dpi=200) if can_render() else []
        for s_idx, slide in enumerate(prs.slides):
            page = Image.open(renders[s_idx]).convert("RGBA") if renders else None
            for n, shape in enumerate(slide.shapes, 1):
                if shape.is_placeholder or not shape.width or not shape.height:
                    continue
                name = re.sub(r"[^\w.-]+", "-", shape.name).strip("-") or f"shape{n}"
                size_in = [emu_to_in(shape.width), emu_to_in(shape.height)]
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    ext = shape.image.ext
                    self.add(shape.image.blob, f"icon-deck/{deck_path.stem}/s{s_idx + 1}-{name}.{ext}", "icon-deck",
                             size_in=size_in)
                    continue
                el = shape._element
                if el.xpath(".//a:t[normalize-space()]") or el.xpath(".//@r:embed | .//@r:link | .//@r:id"):
                    continue  # text labels, and shapes whose parts would not survive a clone
                if page is None:
                    continue  # a vector shape cannot be previewed without a renderer
                sx, sy = page.width / prs.slide_width, page.height / prs.slide_height
                crop = page.crop((int(shape.left * sx), int(shape.top * sy),
                                  int((shape.left + shape.width) * sx) + 1, int((shape.top + shape.height) * sy) + 1))
                crop.thumbnail((256, 256))
                xml = etree.tostring(copy.deepcopy(el))
                self.add(xml, f"icon-deck/{deck_path.stem}/s{s_idx + 1}-{name}.xml", "icon-deck",
                         shape_xml=True, size_in=size_in, preview=crop)
        shutil.rmtree(self.previews / "_deck", ignore_errors=True)

    def finish(self):
        # Assets from folders/decks not passed on this run stay registered while their files exist.
        for digest, row in self.cache.items():
            if digest not in self.assets and row["source"] != "template" and (self.tdir / row["file"]).exists():
                self.assets[digest] = row
        order = {"logo": 0, "icon": 1, "diagram": 2, "pattern": 3, "photo": 4}
        rows = sorted(self.assets.values(), key=lambda a: (order.get(a["kind"], 9), a["file"]))
        for i, row in enumerate(rows, 1):
            row["index"] = i
        thumbs = self.tdir / "thumbs"
        thumbs.mkdir(exist_ok=True)
        for old in list(thumbs.glob("icons-*.jpg")) + list(thumbs.glob("graphics-*.jpg")):
            old.unlink()
        for prefix, kinds in (("icons", {"icon", "logo"}), ("graphics", {"diagram", "pattern", "photo"})):
            items = [(f"#{a['index']} {a['label']}", self.tdir / a["preview"])
                     for a in rows if a["kind"] in kinds and a["preview"]]
            sheets = contact_sheets(items, thumbs / prefix, cols=6, cell_w=160, per_sheet=48, checker=True)
            shown = [a for a in rows if a["kind"] in kinds and a["preview"]]
            for i, a in enumerate(shown):
                a["thumb"] = f"thumbs/{sheets[i // 48].name}"
        write_json(self.tdir / "assets.json", {"assets": rows})
        return rows


def build_assets(tdir, icons=(), graphics=(), icon_decks=()):
    h = Harvest(tdir)
    h.template_media()
    for d in icons:
        h.folder(d, "icons")
    for d in graphics:
        h.folder(d, "graphics")
    for deck in icon_decks:
        h.icon_deck(deck)
    return h.finish()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template_dir")
    ap.add_argument("--icons", action="append", default=[])
    ap.add_argument("--graphics", action="append", default=[])
    ap.add_argument("--icon-deck", action="append", default=[])
    args = ap.parse_args()
    rows = build_assets(args.template_dir, args.icons, args.graphics, args.icon_deck)
    unlabeled = sum(1 for a in rows if not a["labeled"])
    print(f"wrote assets.json: {len(rows)} assets ({unlabeled} with filename-guessed labels awaiting review)")


if __name__ == "__main__":
    main()
